"""Targeted OOXML edits; unrelated ZIP members remain byte-for-byte unchanged."""
import hashlib
import io
import json
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from workbench.files import MAX_FILE_BYTES

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
S = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def archive(data):
    z = zipfile.ZipFile(io.BytesIO(data))
    names = z.namelist()
    if len(names) != len(set(names)) or len(names) > 10000 or sum(x.file_size for x in z.infolist()) > MAX_FILE_BYTES:
        z.close()
        raise ValueError('Office 压缩包重复条目、条目过多或解压后超过 200 MB')
    if any('_xmlsignatures/' in n or n.endswith('vbaProject.bin') for n in names):
        z.close()
        raise ValueError('暂不编辑带数字签名或宏的文档')
    return z


def xml(data):
    if b'<!DOCTYPE' in data or b'<!ENTITY' in data:
        raise ValueError('不支持带 DTD/实体的 XML')
    for _, pair in ET.iterparse(io.BytesIO(data), events=['start-ns']):
        prefix, uri = pair
        if not re.fullmatch(r'ns\d+', prefix):
            ET.register_namespace(prefix, uri)
    return ET.fromstring(data)


def serialize(root, source):
    # Preserve namespace declarations referenced by mc:Ignorable and QName values.
    from xml.sax.saxutils import quoteattr
    encoded = ET.tostring(root, encoding='unicode')
    head = encoded.split('>', 1)[0]
    declarations = {}
    for _, (prefix, uri) in ET.iterparse(io.BytesIO(source), events=['start-ns']):
        declarations.setdefault(prefix, uri)
    for prefix, uri in declarations.items():
        attr = 'xmlns:' + prefix if prefix else 'xmlns'
        if not re.search(r'\s' + re.escape(attr) + r'=', head):
            root.set(attr, uri)
    return ET.tostring(root, encoding='utf-8', xml_declaration=True)


def targets(z, suffix):
    if suffix == '.docx':
        parts = ['word/document.xml']
        namespace = W
    else:
        parts = sorted((n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)), key=lambda n: int(re.search(r'(\d+)\.xml', n).group(1)))
        namespace = A
    rows = []
    for part in parts:
        root = xml(z.read(part))
        for i, paragraph in enumerate(root.iter('{%s}p' % namespace), 1):
            text = ''.join(t.text or '' for t in paragraph.iter('{%s}t' % namespace))
            rows.append({'target': part + '#p' + str(i), 'text': text})
    return rows


def sheets(z):
    import posixpath
    rels = {x.get('Id'): x.get('Target') for x in xml(z.read('xl/_rels/workbook.xml.rels')) if x.get('TargetMode') != 'External'}
    result = {}
    for sheet in xml(z.read('xl/workbook.xml')).iter('{%s}sheet' % S):
        target = rels[sheet.get('{%s}id' % R)]
        part = posixpath.normpath(target.lstrip('/') if target.startswith('/') else 'xl/' + target)
        if not part.startswith('xl/worksheets/'):
            raise ValueError('不支持此工作表类型')
        result[sheet.get('name')] = part
    return result


def cell_value(cell, shared):
    if cell is None: return None
    formula = cell.find('{%s}f' % S)
    if formula is not None: return '=' + (formula.text or '')
    if cell.get('t') == 'inlineStr': return ''.join(t.text or '' for t in cell.iter('{%s}t' % S))
    value = cell.findtext('{%s}v' % S)
    if cell.get('t') == 's': return shared[int(value)]
    if cell.get('t') == 'b': return value == '1'
    if value is None: return None
    if cell.get('t') in ('str', 'e'): return value
    try: return float(value)
    except ValueError: return value


def inspect_office(ws, path, start=1, count=100, sheet='', cells=None):
    if type(start) is not int or start < 1 or type(count) is not int or not 1 <= count <= 300:
        raise ValueError('分页范围错误')
    data = ws.bytes(path, MAX_FILE_BYTES)
    suffix = Path(path).suffix.lower()
    with archive(data) as z:
        if suffix in ('.docx', '.pptx'):
            rows = targets(z, suffix)
            result = {'items': rows[start-1:start-1+count], 'total': len(rows)}
        elif suffix == '.xlsx':
            mapping = sheets(z)
            result = {'sheets': list(mapping), 'items': []}
            if sheet:
                if sheet not in mapping: raise ValueError('工作表不存在')
                if not isinstance(cells, list) or not 1 <= len(cells) <= 300: raise ValueError('请指定 1 至 300 个单元格地址')
                shared = []
                if 'xl/sharedStrings.xml' in z.namelist():
                    shared = [''.join(t.text or '' for t in item.iter('{%s}t' % S)) for item in xml(z.read('xl/sharedStrings.xml'))]
                root = xml(z.read(mapping[sheet]))
                lookup = {c.get('r'): c for c in root.iter('{%s}c' % S)}
                result['items'] = [{'cell': c, 'value': cell_value(lookup.get(c), shared)} for c in cells]
        else: raise ValueError('仅支持 docx、pptx、xlsx')
    result.update(path=path, sha256=hashlib.sha256(data).hexdigest())
    return result


def edit_office(ws, path, expected_sha256, edits, output=''):
    if not isinstance(edits, list) or not edits: raise ValueError('edits 必须为非空列表')
    suffix = Path(path).suffix.lower()
    if suffix not in ('.docx', '.pptx', '.xlsx'): raise ValueError('仅支持 docx、pptx、xlsx')
    output = output or path
    if Path(output).suffix.lower() != suffix: raise ValueError('另存为必须保持原格式')
    original = ws.bytes(path, MAX_FILE_BYTES)
    if hashlib.sha256(original).hexdigest() != expected_sha256: raise ValueError('文件版本变化，请重新检查文件')
    with archive(original) as z:
        changes = {}
        checks = []
        for edit in edits:
            if not isinstance(edit, dict): raise ValueError('每项修改必须是对象')
            required = {'target','old_text','text'} if suffix in ('.docx','.pptx') else {'sheet','cell','value'}
            if not required.issubset(edit): raise ValueError('修改参数缺少字段：' + ', '.join(sorted(required-set(edit))))
            if suffix in ('.docx', '.pptx'):
                target = edit['target']
                part, index = target.rsplit('#p', 1)
                allowed = {row['target'] for row in targets(z, suffix)}
                if target not in allowed: raise ValueError('段落定位不存在')
                root = xml(changes.get(part, z.read(part)))
                ns = W if suffix == '.docx' else A
                paragraph = list(root.iter('{%s}p' % ns))[int(index)-1]
                texts = list(paragraph.iter('{%s}t' % ns))
                if ''.join(t.text or '' for t in texts) != edit['old_text']: raise ValueError('原段落文本不匹配')
                if not isinstance(edit['text'], str) or any(c in edit['text'] for c in '\n\r\t'): raise ValueError('段落替换仅接受单段纯文本')
                if not texts: raise ValueError('空段落暂不支持；请选择已有文本段落')
                # Preserve runs and formatting; replacement inherits first text run style.
                texts[0].text = edit['text']
                texts[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                for node in texts[1:]: node.text = ''
                checks.append((target, edit['text']))
            else:
                mapping = sheets(z)
                if edit['sheet'] not in mapping: raise ValueError('工作表不存在')
                part = mapping[edit['sheet']]
                root = xml(changes.get(part, z.read(part)))
                address = edit['cell']
                match = re.fullmatch(r'([A-Z]{1,3})([1-9][0-9]{0,6})', address)
                if not match: raise ValueError('单元格地址无效')
                col = 0
                for char in match[1]: col = col*26 + ord(char)-64
                if col > 16384 or int(match[2]) > 1048576: raise ValueError('单元格越界')
                cell = next((c for c in root.iter('{%s}c' % S) if c.get('r') == address), None)
                if cell is None: raise ValueError('首版仅修改文件中已有的单元格，请先检查定位')
                if cell.find('{%s}f' % S) is not None and cell.find('{%s}f' % S).attrib:
                    raise ValueError('不支持修改共享或数组公式')
                value = edit['value']
                for child in list(cell):
                    if child.tag in ('{%s}v' % S, '{%s}f' % S, '{%s}is' % S): cell.remove(child)
                cell.attrib.pop('t', None)
                if isinstance(value, bool):
                    cell.set('t','b'); ET.SubElement(cell,'{%s}v' % S).text = '1' if value else '0'
                elif isinstance(value, (float,int)):
                    import math
                    if not math.isfinite(value): raise ValueError('数值必须有限')
                    ET.SubElement(cell,'{%s}v' % S).text = str(value)
                elif isinstance(value,str):
                    if len(value) > 32767: raise ValueError('单元格文字过长')
                    if value.startswith('='): raise ValueError('首版不写入公式；避免计算链和缓存不一致')
                    cell.set('t','inlineStr'); ET.SubElement(ET.SubElement(cell,'{%s}is' % S),'{%s}t' % S).text=value
                elif value is not None: raise ValueError('单元格仅接受文本、数字、布尔值或 null')
                checks.append((edit['sheet'],address,value))
            changes[part] = serialize(root, changes.get(part, z.read(part)))
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer,'w',zipfile.ZIP_DEFLATED) as dest:
            for entry in z.infolist(): dest.writestr(entry, changes.get(entry.filename,z.read(entry.filename)))
    data = buffer.getvalue()
    if len(data)>MAX_FILE_BYTES: raise ValueError('输出超过 200 MB')
    def verify(payload):
        with archive(payload) as saved:
            if saved.testzip(): raise ValueError('Office ZIP 完整性验证失败')
            if suffix in ('.docx','.pptx'):
                values = {r['target']:r['text'] for r in targets(saved,suffix)}
                for target,value in checks:
                    if values[target]!=value: raise ValueError('段落保存验证失败')
            else:
                for sheet,address,value in checks:
                    root = xml(saved.read(sheets(saved)[sheet]))
                    cell = next(c for c in root.iter('{%s}c' % S) if c.get('r')==address)
                    if cell_value(cell,[])!=value: raise ValueError('单元格保存验证失败')
    verify(data)
    backup = None
    import uuid, os
    with ws.write_target(output) as target:
        if output != path and target.exists(): raise ValueError('另存为目标已存在，请选择新文件名')
        if ws.bytes(path,MAX_FILE_BYTES)!=original: raise ValueError('源文件变化，已停止保存')
        if output==path:
            backup_name=target.name+'.'+uuid.uuid4().hex+'.bak'
            with (target.parent/backup_name).open('xb') as f: f.write(original)
            backup=(Path(output).parent/backup_name).as_posix()
        temp=target.parent/('.office-'+uuid.uuid4().hex)
        try:
            with temp.open('xb') as f:
                f.write(data);f.flush();os.fsync(f.fileno())
            verify(temp.read_bytes())
            if output==path: os.replace(temp,target)
            else:
                # Exclusive destination creation prevents overwriting another user's file.
                with target.open('xb') as f: f.write(data)
        finally:
            if temp.exists(): temp.unlink()
    verify(ws.bytes(output,MAX_FILE_BYTES))
    return {'path':output,'backup':backup,'sha256':hashlib.sha256(data).hexdigest(),'verified':True,'changes':len(edits),'note':'结构及修改值已重新读取验证；未验证视觉排版。段落替换继承首个文本片段格式；Excel 公式结果需在 Office 中重算。'}


def create_office(ws, path, content):
    """Create a new Office file, reopen and verify content before/after exclusive save."""
    suffix = Path(path).suffix.lower()
    if not isinstance(content, dict): raise ValueError('content 必须为对象')
    def text(value):
        if not isinstance(value, str): raise ValueError('标题和段落必须为字符串')
        return value
    buffer = io.BytesIO()
    try:
        if suffix == '.docx':
            from docx import Document
            doc = Document()
            title = text(content.get('title', ''))
            paragraphs = content.get('paragraphs', [])
            tables = content.get('tables', [])
            if not isinstance(paragraphs, list) or not isinstance(tables, list): raise ValueError('paragraphs/tables 必须为列表')
            expected = []
            if title: doc.add_heading(title, 0); expected.append(title)
            for paragraph in paragraphs:
                paragraph = text(paragraph); doc.add_paragraph(paragraph); expected.append(paragraph)
            for rows in tables:
                if not isinstance(rows, list) or not rows or not all(isinstance(row,list) for row in rows): raise ValueError('表格必须为非空二维字符串数组')
                width = len(rows[0])
                if width == 0 or any(len(row)!=width for row in rows): raise ValueError('表格列数必须一致')
                table = doc.add_table(rows=len(rows), cols=width)
                table.style = 'Table Grid'
                for i,row in enumerate(rows):
                    for j,value in enumerate(row): table.cell(i,j).text = text(value)
            doc.save(buffer)
            def verify(data):
                saved = Document(io.BytesIO(data))
                if [p.text for p in saved.paragraphs] != expected: raise ValueError('Word 段落验证失败')
                actual = [[[c.text for c in row.cells] for row in table.rows] for table in saved.tables]
                if actual != tables: raise ValueError('Word 表格验证失败')
        elif suffix == '.xlsx':
            from openpyxl import Workbook, load_workbook
            from openpyxl.styles import Font
            import math
            sheets_spec = content.get('sheets', [{'name':'Sheet1','rows':[]}])
            if not isinstance(sheets_spec,list) or not sheets_spec: raise ValueError('至少提供一个工作表')
            book = Workbook(); book.remove(book.active)
            names = set(); expected = []
            for spec in sheets_spec:
                name = text(spec['name'])
                if not name or len(name)>31 or re.search(r'[\\/*?:\[\]]',name) or name.casefold() in names: raise ValueError('工作表名称无效或重复')
                names.add(name.casefold()); sheet = book.create_sheet(name)
                rows = spec.get('rows',[])
                if not isinstance(rows,list) or len(rows)>1048576: raise ValueError('rows 必须为合法行数组')
                for i,row in enumerate(rows,1):
                    if not isinstance(row,list) or len(row)>16384: raise ValueError('每行必须为合法数组')
                    for j,value in enumerate(row,1):
                        if value is not None and not isinstance(value,(str,int,float,bool)): raise ValueError('单元格类型不支持')
                        if isinstance(value,float) and not math.isfinite(value): raise ValueError('数字必须有限')
                        if isinstance(value,str) and len(value)>32767: raise ValueError('单元格文字过长')
                        cell=sheet.cell(i,j,value)
                        # Input strings are literal; never silently turn user content into formulas.
                        if isinstance(value,str): cell.data_type='s'
                        expected.append((name,i,j,None if value=='' else value))
                if rows:
                    for cell in sheet[1]: cell.font=Font(bold=True)
                    sheet.freeze_panes='A2'
            book.save(buffer);book.close()
            def verify(data):
                saved=load_workbook(io.BytesIO(data),data_only=False)
                try:
                    if saved.sheetnames != [s['name'] for s in sheets_spec]: raise ValueError('工作表验证失败')
                    for name,i,j,value in expected:
                        if saved[name].cell(i,j).value != value: raise ValueError('Excel 单元格验证失败')
                finally: saved.close()
        elif suffix == '.pptx':
            from pptx import Presentation
            slides = content.get('slides',[])
            if not isinstance(slides,list) or not slides: raise ValueError('slides 必须包含至少一页')
            deck=Presentation(); expected=[]
            for spec in slides:
                title=text(spec.get('title','')); bullets=spec.get('bullets',[])
                if not isinstance(bullets,list): raise ValueError('bullets 必须为字符串数组')
                slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text=title
                frame=slide.placeholders[1].text_frame;frame.clear()
                for i,value in enumerate(bullets):
                    paragraph=frame.paragraphs[0] if i==0 else frame.add_paragraph()
                    paragraph.text=text(value)
                expected.append((title,'\n'.join(bullets)))
            deck.save(buffer)
            def verify(data):
                saved=Presentation(io.BytesIO(data))
                actual=[(s.shapes.title.text,s.placeholders[1].text) for s in saved.slides]
                if actual!=expected: raise ValueError('PPT 幻灯片验证失败')
        else: raise ValueError('仅支持新建 .docx、.xlsx、.pptx')
    except ImportError as exc:
        raise ValueError('当前解释器缺少创建依赖，请在环境配置中从 NAS 安装 python-docx、openpyxl、python-pptx 及所需依赖') from exc
    except (KeyError, TypeError) as exc:
        raise ValueError('内容字段缺失或类型错误：'+str(exc)) from exc
    data=buffer.getvalue()
    if len(data)>MAX_FILE_BYTES: raise ValueError('生成文件超过 200 MB')
    verify(data)
    with ws.write_target(path) as target:
        with target.open('xb') as stream: stream.write(data)
    saved=ws.bytes(path,MAX_FILE_BYTES)
    verify(saved)
    return {'path':path,'bytes':len(saved),'sha256':hashlib.sha256(saved).hexdigest(),'verified':True,
            'note':'新文件已创建并重新读取验证内容；尚未进行视觉排版验证。'}
