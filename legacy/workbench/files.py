"""有边界、有限量的工作区文件工具，文本覆盖前保留备份。"""
import contextlib
import hashlib
import io
import os
from pathlib import Path
import stat
import zipfile
import uuid

MAX_FILE_BYTES = 200 * 1024 * 1024

TEXT = {'.txt', '.md', '.csv', '.tsv', '.json', '.yaml', '.yml', '.py', '.js', '.ts', '.html',
        '.css', '.xml', '.ini', '.cfg', '.log', '.sql', '.sh', '.rst', '.toml'}
DOCUMENTS = {'.xlsx', '.docx', '.pdf', '.pptx'}
SKIP = {'.git', '.venv', 'env', 'runtime', 'node_modules', '__pycache__', '.test-kylin'}


class Workspace:
    def __init__(self, root):
        self.root = Path(root).resolve(strict=True)
        if not self.root.is_dir():
            raise ValueError('工作目录不存在')

    @contextlib.contextmanager
    def write_target(self, relative):
        requested = Path(relative)
        if (not relative or requested.is_absolute() or requested.drive or
                any(p in ('.', '..') or p.startswith('.') or ':' in p or p in SKIP
                    for p in requested.parts) or not requested.parts):
            raise ValueError('写入必须使用工作目录内的普通相对路径，不支持隐藏目录或环境目录')
        parent = self.root
        descriptors = []
        try:
            if os.name == 'posix':
                flags = os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW
                descriptors.append(os.open(parent, flags))
                for part in requested.parts[:-1]:
                    descriptors.append(os.open(part, flags, dir_fd=descriptors[-1]))
                parent = Path('/proc/self/fd') / str(descriptors[-1])
                parent.resolve(strict=True).relative_to(self.root)
            else:
                for part in requested.parts[:-1]:
                    parent = parent / part
                    if parent.is_symlink() or getattr(parent, 'is_junction', lambda: False)():
                        raise ValueError('写入路径不允许经过链接目录')
                parent.resolve(strict=True).relative_to(self.root)
            target = parent / requested.name
            if target.is_symlink() or getattr(target, 'is_junction', lambda: False)():
                raise ValueError('不允许写入符号链接')
            yield target
        finally:
            for descriptor in reversed(descriptors):
                os.close(descriptor)

    def create_directory(self, path):
        with self.write_target(path) as target:
            target.mkdir()  # 父目录须已存在；不覆盖已有条目。
        return {'path': path, 'created': True}

    def write_text(self, path, content, expected_sha256=''):
        if Path(path).suffix.lower() not in TEXT:
            raise ValueError('写入仅支持文本、代码和配置格式，不支持 Office/PDF 或二进制')
        if not isinstance(content, str) or '\x00' in content:
            raise ValueError('内容必须为普通文本')
        data = content.encode('utf-8')
        if len(data) > MAX_FILE_BYTES:
            raise ValueError('文本写入上限为 200 MB')
        backup = None
        with self.write_target(path) as target:
            if target.exists():
                old = self.bytes(path, MAX_FILE_BYTES)
                if not expected_sha256 or hashlib.sha256(old).hexdigest() != expected_sha256:
                    raise ValueError('覆盖前必须读取文件，并提供最新 sha256；文件可能已改变')
                backup_name = target.name + '.' + uuid.uuid4().hex + '.bak'
                with (target.parent / backup_name).open('xb') as stream:
                    stream.write(old)
                backup = (Path(path).parent / backup_name).as_posix()
                temporary = target.parent / ('.write-' + uuid.uuid4().hex)
                try:
                    with temporary.open('xb') as stream:
                        stream.write(data)
                        stream.flush()
                        os.fsync(stream.fileno())
                    if self.bytes(path, MAX_FILE_BYTES) != old:
                        raise ValueError('写入前文件发生变化，已停止覆盖')
                    os.replace(temporary, target)
                finally:
                    if temporary.exists(): temporary.unlink()
            else:
                if expected_sha256:
                    raise ValueError('原文件已不存在，已停止覆盖')
                with target.open('xb') as stream:
                    stream.write(data)
        return {'path': path, 'bytes': len(data), 'sha256': hashlib.sha256(data).hexdigest(),
                'backup': backup, 'encoding': 'UTF-8'}

    def resolve(self, relative='.'):
        requested = Path(relative)
        if requested.is_absolute() or '..' in requested.parts:
            raise ValueError('请使用工作目录内的相对路径，不允许 ../')
        path = (self.root / requested).resolve(strict=True)
        try:
            path.relative_to(self.root)
        except ValueError:
            raise ValueError('路径或符号链接指向工作目录之外') from None
        return path

    def list(self, relative='.'):
        folder = self.resolve(relative)
        if not folder.is_dir():
            raise ValueError('此路径不是目录')
        items = []
        # 使用有上限的迭代，避免巨大目录阻塞界面。
        truncated = False
        visited = 0
        with os.scandir(folder) as entries:
            for entry in entries:
                visited += 1
                if len(items) >= 300 or visited > 2000:
                    truncated = True
                    break
                if entry.name in SKIP or entry.name.startswith('.'):
                    continue
                logical = (Path(relative) / entry.name).as_posix()
                try:
                    path = self.resolve(logical)
                    info = path.stat()
                    if not (stat.S_ISDIR(info.st_mode) or stat.S_ISREG(info.st_mode)):
                        continue
                    items.append({'name': entry.name, 'path': logical, 'directory': path.is_dir(),
                                  'bytes': info.st_size, 'type': path.suffix.lower(), 'modified': info.st_mtime})
                except (OSError, ValueError):
                    continue
        items.sort(key=lambda x: (not x['directory'], x['name'].casefold()))
        return {'path': relative, 'items': items, 'truncated': truncated}

    def bytes(self, relative, limit):
        path = self.resolve(relative)
        # 非阻塞打开，随后拒绝设备/FIFO；Linux 再根据真实句柄复核目录边界。
        descriptor = os.open(path, os.O_RDONLY | getattr(os, 'O_NONBLOCK', 0) | getattr(os, 'O_NOFOLLOW', 0))
        with os.fdopen(descriptor, 'rb') as source:
            info = os.fstat(source.fileno())
            if not stat.S_ISREG(info.st_mode):
                raise ValueError('只支持普通文件')
            if info.st_size > limit:
                raise ValueError('文件超过本次读取上限（{} MB）'.format(limit // 1024 // 1024))
            handle = Path('/proc/self/fd') / str(source.fileno())
            if handle.exists():
                try:
                    handle.resolve().relative_to(self.root)
                except ValueError:
                    raise ValueError('文件真实位置已越出工作目录') from None
            data = source.read(limit + 1)
            if len(data) > limit:
                raise ValueError('文件在读取时超过大小上限')
            return data

    def read(self, path, start=1, count=120):
        if type(start) is not int or start < 1 or type(count) is not int or not 1 <= count <= 300:
            raise ValueError('起始行需大于 0，每次最多读取 300 行')
        suffix = self.resolve(path).suffix.lower()
        if suffix not in TEXT | DOCUMENTS:
            raise ValueError('暂不支持该格式；支持文本、.xlsx、.docx、.pdf、.pptx（旧版 .doc/.xls/.ppt 请先转换）')
        data = self.bytes(path, MAX_FILE_BYTES)
        extracted_truncated = False
        if suffix in TEXT:
            text = None
            for encoding in ('utf-8-sig', 'gb18030'):
                try:
                    text = data.decode(encoding)
                    break
                except UnicodeDecodeError:
                    pass
            if text is None or '\x00' in text:
                raise ValueError('不是支持的文本编码（UTF-8/GB18030），或文件为二进制')
            lines = text.splitlines()
        else:
            lines, extracted_truncated = document_lines(suffix, data)
        selected = lines[start - 1:start - 1 + count]
        content = '\n'.join('{}: {}'.format(i, line[:2000]) for i, line in enumerate(selected, start))
        return {'path': path, 'sha256': hashlib.sha256(data).hexdigest(), 'start': start, 'total_lines': len(lines), 'content': content[:24000],
                'truncated': extracted_truncated or start - 1 + count < len(lines) or len(content) > 24000
                             or any(len(x) > 2000 for x in selected),
                'note': '文档行号为解析后的摘录行；段落、页码、工作表或幻灯片位置见内容标注。'
                        if suffix in DOCUMENTS else '行号对应原文本。'}

    def search(self, query, relative='.'):
        if not isinstance(query, str) or not query.strip() or len(query) > 200:
            raise ValueError('搜索词长度需为 1 到 200 字符')
        folder = self.resolve(relative)
        if not folder.is_dir():
            raise ValueError('搜索范围必须为目录')
        scanned = skipped = 0
        matches = []
        truncated = False
        visited_dirs = visited_entries = 0
        for current, dirs, names in os.walk(folder, followlinks=False):
            visited_dirs += 1
            if visited_dirs > 200:
                truncated = True
                break
            dirs[:] = sorted(d for d in dirs if d not in SKIP and not d.startswith('.') and not (Path(current)/d).is_symlink())
            for name in sorted(names):
                visited_entries += 1
                if scanned >= 600 or len(matches) >= 60 or visited_entries > 3000:
                    truncated = True
                    break
                path = (Path(current) / name).relative_to(self.root).as_posix()
                if Path(name).suffix.lower() not in TEXT:
                    skipped += 1
                    continue
                scanned += 1
                try:
                    data = self.bytes(path, MAX_FILE_BYTES)
                    try:
                        text = data.decode('utf-8-sig')
                    except UnicodeDecodeError:
                        text = data.decode('gb18030')
                    if '\x00' in text:
                        continue
                    for number, line in enumerate(text.splitlines(), 1):
                        if query.casefold() in line.casefold():
                            matches.append({'path': path, 'line': number, 'text': line[:400]})
                            if len(matches) >= 60:
                                truncated = True
                                break
                except (OSError, ValueError, UnicodeError):
                    skipped += 1
            if truncated:
                break
        return {'query': query, 'matches': matches, 'scanned': scanned, 'skipped': skipped,
                'truncated': truncated, 'scope': '搜索文本文件；Office/PDF 请通过读取文档工具查看。'}


def document_lines(suffix, data):
    if suffix in {'.xlsx', '.docx', '.pptx'}:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 10000 or sum(m.file_size for m in entries) > MAX_FILE_BYTES:
                raise ValueError('文档解压后过大，请拆分后读取')
    lines = []
    truncated = False
    def add(label, text):
        nonlocal truncated
        if len(lines) >= 2000:
            truncated = True
            return False
        value = str(text)
        if len(value) > 2000:
            truncated = True
        lines.append(label + value[:2000])
        return True
    try:
        if suffix == '.docx':
            from docx import Document
            doc = Document(io.BytesIO(data))
            for i, paragraph in enumerate(doc.paragraphs, 1):
                if not add('[段落 {}] '.format(i), paragraph.text): break
            for i, table in enumerate(doc.tables, 1):
                for j, row in enumerate(table.rows, 1):
                    if not add('[表 {} 行 {}] '.format(i, j), ' | '.join(c.text for c in row.cells)): break
                if truncated and len(lines) >= 2000: break
        elif suffix == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in book.worksheets[:20]:
                    for number, row in enumerate(sheet.iter_rows(max_row=500, max_col=50, values_only=True), 1):
                        if any(v is not None for v in row):
                            if not add('[{} 行 {}] '.format(sheet.title, number), ' | '.join('' if v is None else str(v) for v in row)): break
                    truncated |= (sheet.max_row or 0) > 500 or (sheet.max_column or 0) > 50
                    if len(lines) >= 2000: break
                truncated |= len(book.worksheets) > 20
                add('[说明] ', '公式只读取文件中缓存的结果，不执行公式。')
            finally:
                book.close()
        elif suffix == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted: raise ValueError('不支持加密 PDF')
            total = len(reader.pages)
            truncated |= total > 30
            for number in range(min(total, 30)):
                text = reader.pages[number].extract_text() or ''
                for line in text.splitlines() or ['（无可提取文本，扫描页需要另行 OCR）']:
                    if not add('[第 {} 页] '.format(number + 1), line): break
                if len(lines) >= 2000: break
        else:
            from pptx import Presentation
            deck = Presentation(io.BytesIO(data))
            truncated |= len(deck.slides) > 100
            for number, slide in enumerate(deck.slides, 1):
                if number > 100: break
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        if not add('[幻灯片 {}] '.format(number), shape.text): break
                    elif shape.has_table:
                        for row in shape.table.rows:
                            if not add('[幻灯片 {} 表格] '.format(number), ' | '.join(c.text for c in row.cells)): break
                    if len(lines) >= 2000: break
                if len(lines) >= 2000: break
    except ImportError as exc:
        raise ValueError('当前解释器缺少文档依赖。请在环境配置中从 NAS 为当前解释器安装 openpyxl、python-docx、pypdf、python-pptx 及所需依赖。') from exc
    return lines, truncated
